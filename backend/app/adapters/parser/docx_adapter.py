"""DOCX text extraction adapter using python-docx."""

import io
from typing import List
import docx
from docx.opc.exceptions import PackageNotFoundError

from app.adapters.parser.base import ExtractedDocument, PageSegment
from app.adapters.parser.validator import ParserException


def extract_docx(content: bytes) -> ExtractedDocument:
    """Extract text from DOCX paragraphs and tables."""
    segments: List[PageSegment] = []
    warnings: List[str] = []
    full_text_parts: List[str] = []

    try:
        stream = io.BytesIO(content)
        doc = docx.Document(stream)

        # Extract paragraphs
        for idx, paragraph in enumerate(doc.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                segments.append(PageSegment(page_number=1, text=text))
                full_text_parts.append(text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    combined = " | ".join(row_texts)
                    segments.append(PageSegment(page_number=1, text=combined))
                    full_text_parts.append(combined)

    except (PackageNotFoundError, Exception) as e:
        raise ParserException(
            "Unable to read or parse DOCX document: malformed or corrupt file structure",
            code="corrupt_file",
            status_code=422,
        )

    raw_text = "\n".join(full_text_parts)
    if not raw_text.strip():
        warnings.append("No extractable text found in DOCX document")

    return ExtractedDocument(
        raw_text=raw_text,
        segments=segments,
        warnings=warnings,
        format_type="docx",
    )
